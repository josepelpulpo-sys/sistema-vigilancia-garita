import datetime
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
import streamlit as st
from sqlalchemy import create_engine, text

# Configuración de la página web
st.set_page_config(
    page_title="Control de Portería & Recepción de Telas",
    layout="wide",
    page_icon="🚛",
)

# -------------------------------------------------------------------
# USUARIOS AUTORIZADOS (ID / CONTRASEÑA / NOMBRE)
# -------------------------------------------------------------------
USUARIOS_AUTORIZADOS = {
    "vigilante1": {"password": "123", "nombre": "Juan Pérez"},
    "vigilante2": {"password": "456", "nombre": "Carlos Gómez"},
    "admin": {"password": "admin", "nombre": "Administrador"}
}


# -------------------------------------------------------------------
# BASE DE DATOS SUPABASE / POSTGRESQL (VÍA SQLALCHEMY)
# -------------------------------------------------------------------
def get_engine():
    db_url = st.secrets["postgres"]["url"]
    return create_engine(db_url)


def inicializar_db():
    engine = get_engine()
    with engine.begin() as conn:
        # 1. Tabla de Movimientos / Bitácora Completa de Eventos
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS historial_movimientos (
                id SERIAL PRIMARY KEY,
                fecha_evento VARCHAR(20),
                hora_evento VARCHAR(20),
                tipo_operacion VARCHAR(50),
                num_guia_entrada VARCHAR(50),
                num_guia_salida VARCHAR(50) DEFAULT '-',
                num_partida VARCHAR(50) DEFAULT '-',
                cliente VARCHAR(100),
                tipo_tela VARCHAR(100),
                rollos_movimiento INT,
                peso_kg_movimiento DOUBLE PRECISION,
                rollos_saldo_planta INT,
                chofer VARCHAR(100),
                dni_chofer VARCHAR(30),
                empresa_transporte VARCHAR(100),
                placa VARCHAR(20),
                vigilante VARCHAR(100)
            )
        """))

        # 2. Tabla de Stock Actual / Estado de Guías en Planta
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS stock_telas_planta (
                id SERIAL PRIMARY KEY,
                num_guia_entrada VARCHAR(50),
                num_partida VARCHAR(50),
                cliente VARCHAR(100),
                tipo_tela VARCHAR(100),
                rollos_actuales INT,
                peso_kg_inicial DOUBLE PRECISION,
                fecha_ingreso VARCHAR(20),
                chofer_ingreso VARCHAR(100),
                dni_chofer_ingreso VARCHAR(30),
                transporte_ingreso VARCHAR(100),
                placa_ingreso VARCHAR(20)
            )
        """))


try:
    inicializar_db()
except Exception as e:
    st.error(f"Error al conectar con la base de datos de Supabase: {e}")


# -------------------------------------------------------------------
# SESIÓN DE OPERADOR / VIGILANTE (LOGIN CON ID Y CONTRASEÑA)
# -------------------------------------------------------------------
if "vigilante" not in st.session_state:
    st.session_state["vigilante"] = None

if not st.session_state["vigilante"]:
    st.markdown("## 🛡️ Control de Portería & Recepción de Telas")
    st.info("Por favor, ingrese sus credenciales de acceso para iniciar sesión.")

    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        with st.form(key="form_login"):
            user_input = st.text_input("ID de Usuario:").strip()
            pass_input = st.text_input("Contraseña:", type="password").strip()
            submit_login = st.form_submit_button("Ingresar al Sistema", use_container_width=True)

            if submit_login:
                if user_input in USUARIOS_AUTORIZADOS and USUARIOS_AUTORIZADOS[user_input]["password"] == pass_input:
                    st.session_state["vigilante"] = USUARIOS_AUTORIZADOS[user_input]["nombre"]
                    st.success(f"Bienvenido/a, **{st.session_state['vigilante']}**.")
                    st.rerun()
                else:
                    st.error("⚠️ ID de usuario o contraseña incorrectos.")
    st.stop()


# -------------------------------------------------------------------
# BARRA LATERAL
# -------------------------------------------------------------------
with st.sidebar:
    st.title("👤 Turno Activo")
    st.write(f"**Vigilante:** {st.session_state['vigilante']}")
    st.write(f"**Fecha:** {datetime.datetime.now().strftime('%Y-%m-%d')}")
    st.markdown("---")
    if st.button("🔴 Cerrar Sesión / Cambiar Usuario", use_container_width=True):
        st.session_state["vigilante"] = None
        st.rerun()


# -------------------------------------------------------------------
# GENERACIÓN DE EXCEL COMPATIBLE
# -------------------------------------------------------------------
def generar_excel_profesional():
    engine = get_engine()
    query = """
        SELECT 
            id AS "N° Movimiento",
            fecha_evento AS "Fecha Evento",
            hora_evento AS "Hora Evento",
            tipo_operacion AS "Tipo de Operación",
            num_guia_entrada AS "N° Guía Entrada",
            num_guia_salida AS "N° Guía Salida",
            num_partida AS "N° Partida",
            cliente AS "Cliente",
            tipo_tela AS "Tipo de Tela",
            rollos_movimiento AS "Rollos Movidis / Retirados",
            peso_kg_movimiento AS "Peso Movidi (Kg)",
            rollos_saldo_planta AS "Saldo Restante en Planta",
            chofer AS "Chófer Registrado",
            dni_chofer AS "DNI Chófer",
            empresa_transporte AS "Empresa Transportista",
            placa AS "Placa Vehículo",
            vigilante AS "Vigilante en Turno"
        FROM historial_movimientos
        ORDER BY id DESC
    """
    df = pd.read_sql_query(query, engine)
    df = df.fillna("-")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial Completo"

    navy_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9")
    )
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        ws.append(row)
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            
            if r_idx == 1:
                cell.fill = navy_header
                cell.font = white_bold
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = Font(name="Calibri", size=10)
                if r_idx % 2 == 0:
                    cell.fill = zebra_fill
                
                if c_idx in [1, 2, 3, 4, 5, 6, 7, 13, 14, 16, 17]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [10, 11, 12]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# -------------------------------------------------------------------
# INTERFAZ PRINCIPAL
# -------------------------------------------------------------------
st.title("🚛 Control de Portería - Bitácora Completa de Movimientos")
st.caption(f"Vigilante en turno: **{st.session_state['vigilante']}**")
st.markdown("---")

tab1, tab2, tab3 = st.tabs([
    "📥 1. Registro de Portería (Dejar / Recoger)",
    "📦 2. Stock Actual en Almacén",
    "📋 3. Historial de Todos los Movimientos"
])

# -------------------------------------------------------------------
# PESTAÑA 1: REGISTRO DE MOVIMIENTO
# -------------------------------------------------------------------
with tab1:
    st.subheader("1. Datos Completos del Vehículo, Chófer y Transporte")
    
    col_v1, col_v2, col_v3, col_v4, col_v5 = st.columns([2.5, 2, 2.5, 2, 3])
    with col_v1:
        chofer_nom = st.text_input("Nombre del Chófer *").strip()
    with col_v2:
        dni_chofer_in = st.text_input("DNI / Documento del Chófer *").strip()
    with col_v3:
        empresa_trans = st.text_input("Empresa de Transporte / Tercero *", placeholder="Ej. Servitrans").strip()
    with col_v4:
        placa_veh = st.text_input("Placa del Carro/Camión *").upper().strip()
    with col_v5:
        operacion = st.radio("¿Qué acción realiza el vehículo? *", ["Dejar Tela (Ingreso a Planta)", "Recoger Tela (Despacho de Planta)"])

    st.markdown("---")

    # OPERACIÓN A: DEJAR TELA
    if operacion == "Dejar Tela (Ingreso a Planta)":
        st.subheader("2. Registro de Guía de Entrada y Cliente")
        
        col_g1, col_g2 = st.columns(2)
        with col_g1:
            num_guia_e = st.text_input("N° Guía de Remisión de Entrada *").strip()
        with col_g2:
            cliente_e = st.text_input("Cliente de la Tela *").strip()

        st.markdown("##### 📝 Lista de telas que vienen en esta Guía:")
        
        if "telas_dejar" not in st.session_state:
            st.session_state["telas_dejar"] = []

        with st.form(key="form_add_tela_dejar", clear_on_submit=True):
            col_d1, col_d2, col_d3, col_d4 = st.columns([2, 3, 2, 2])
            with col_d1:
                partida_t = st.text_input("N° Partida (Opcional)", placeholder="Ej. P-102")
            with col_d2:
                tipo_t = st.text_input("Tipo de Tela (ej. Jersey 18/1, Rib Veteado 24/1)")
            with col_d3:
                cant_r = st.number_input("Cantidad de Rollos", min_value=1, step=1)
            with col_d4:
                peso_k = st.number_input("Peso de Entrada (Kg)", min_value=0.1, step=0.1)
            
            btn_add = st.form_submit_button("➕ Agregar Tela a
